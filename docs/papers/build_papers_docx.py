"""네 논문의 Markdown을 같은 판형의 Word 원고로 만든다.

사용: python docs/papers/build_papers_docx.py [--src 원고폴더] [--out 출력폴더] [--stamp YYYYMMDD]
  --src    paper-*.md 가 있는 폴더. 기본값은 docs/papers 에 원고가 있으면 그곳, 없으면 docs/papers/논문 초안
  --out    Word 파일을 둘 폴더. 기본값 docs/papers/output
  --stamp  파일명 날짜. 기본값 20260919 (기존 파일명 규칙을 유지한다)
"""
from pathlib import Path
import argparse
import importlib.util
import json
import os
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location('page_builder', BASE/'page-study/build_page_docx.py')
helper = importlib.util.module_from_spec(spec)
spec.loader.exec_module(helper)
TITLES = {
    'Platform': '거점 사전분포를 고려한 상권 그래프 신경망 업종 분류의 성능 한계',
    'Page': '공공행정자료 결합 기반 상업용 공실 정보의 데이터 품질과 계산 재현성',
    'Posting': '공개 데이터 기반 상가 입점 회수기간 시나리오의 해상도 한계',
    'Program': '창업 검증 프로그램 생성의 근거 계약과 사실성 가드레일',
}
TEXT_WIDTH_CM = 16.6   # A4 21cm - 좌우 여백 2.2cm x 2


def default_src():
    return BASE if (BASE/'paper-platform.md').exists() else BASE/'논문 초안'


def display_target(src_dir, target):
    """원고 폴더 기준 상대 링크를 docs/papers 기준 경로로 바꿔 '자료 주석'에 보인다."""
    path, _, frag = target.partition('#')
    if not path:
        return target
    try:
        rel = Path(os.path.relpath((src_dir/path).resolve(), BASE)).as_posix()
    except ValueError:      # 드라이브가 다른 경우에는 원 표기를 그대로 둔다
        rel = path
    return rel + ('#'+frag if frag else '')


def table_widths(values):
    """열 길이에 비례해 너비를 나눈다. 긴 칸이 판면을 독점하지 않도록 상한을 둔다."""
    def plain(v):
        return re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', v).replace('**', '').replace('`', '')
    count = len(values[0])
    weights = []
    for j in range(count):
        lens = [len(plain(r[j])) for r in values]
        w = max(8, min(sum(lens)/len(lens), 70))
        if values[0][j] == '판정':      # 판정 칸은 두세 줄로 접히지 않게 넉넉히 둔다
            w = max(w, 16)
        weights.append(w)
    raw = [max(2.3, TEXT_WIDTH_CM*x/sum(weights)) for x in weights]
    scale = TEXT_WIDTH_CM/sum(raw)
    return [x*scale for x in raw]


def build(track, title, src_dir, out_dir, stamp):
    src = (src_dir/f'paper-{track.lower()}.md').read_text(encoding='utf-8-sig')
    src = re.sub(r'<!--.*?-->', '', src, flags=re.S)
    src = '## 초록' + src.split('## 초록', 1)[1]
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin, s.bottom_margin = Cm(2.1), Cm(2)
    s.left_margin = s.right_margin = Cm(2.2)
    s.footer_distance = Cm(1)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','Footer'):
        st = doc.styles[name]
        st.font.name = helper.FONT
        st.font.color.rgb = RGBColor(0,0,0)
        rf = st.element.get_or_add_rPr().get_or_add_rFonts()
        for k in ('ascii','hAnsi','eastAsia','cs'): rf.set(qn('w:'+k), helper.FONT)
        for k in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'): rf.attrib.pop(qn('w:'+k),None)
    n = doc.styles['Normal']
    n.font.size = Pt(10.5)
    n.paragraph_format.line_spacing = 1.4
    n.paragraph_format.space_after = Pt(7)
    n.paragraph_format.widow_control = True
    for name,size in [('Title',19),('Heading 1',14),('Heading 2',11.5)]:
        st=doc.styles[name]; st.font.size=Pt(size); st.font.bold=True
        st.paragraph_format.space_before=Pt(13 if name!='Title' else 0)
        st.paragraph_format.space_after=Pt(7)
        st.paragraph_format.keep_with_next=True
    p=doc.add_paragraph(title,'Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph('PlaceOS '+track,'Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.styles['Subtitle'].font.size=Pt(11)
    doc.styles['Subtitle'].paragraph_format.space_after=Pt(14)
    p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)
    doc.core_properties.title=title
    doc.core_properties.subject='PlaceOS '+track+' 연구'
    doc.core_properties.author=''; doc.core_properties.last_modified_by=''
    for name in ('embedTrueTypeFonts','saveSubsetFonts'):
        el=OxmlElement('w:'+name); el.set(qn('w:val'),'true' if name=='embedTrueTypeFonts' else 'false'); doc.settings.element.append(el)
    notes={}

    def inline(p, text):
        pattern=r'(\[[^\]]+\]\([^)]+\)|\*\*.*?\*\*|`[^`]+`|\*[^*]+\*|https?://[^\s]+)'
        for tok in re.split(pattern,text):
            if not tok: continue
            m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',tok)
            if m:
                label,target=m.groups()
                if target.startswith('http'): helper.hyperlink(p,label,target)
                else:
                    shown=display_target(src_dir,target)
                    if shown not in notes: notes[shown]=(len(notes)+1,label)
                    helper.hyperlink(p,label+f' [자료 {notes[shown][0]}]',anchor=f'source{notes[shown][0]}')
            elif tok.startswith('http'): helper.hyperlink(p,tok.rstrip('.'),tok.rstrip('.'))
            else:
                bold=tok.startswith('**'); italic=tok.startswith('*') and not bold
                r=p.add_run(tok.strip('*') if bold or italic else tok.strip('`'))
                helper.font_run(r); r.bold=bold; r.italic=italic

    lines=src.splitlines(); i=0; table_count=0
    while i<len(lines):
        line=lines[i].strip(); i+=1
        if not line or line=='---' or line.startswith('<a '): continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):
                rows.append(lines[i].strip()); i+=1
            rows=[x for x in rows if not re.match(r'^\|[\s:|\-]+\|$',x)]
            values=[[v.strip() for v in x.strip('|').split('|')] for x in rows]
            count=len(values[0]); table_count+=1
            assert all(len(r)==count for r in values), (track, values[0])
            widths=table_widths(values)
            # 열 전체가 짧을 때만 가운데 정렬한다(한 열 안에서 정렬이 섞이지 않게)
            short_col=[all(len(r[j])<20 for r in values) for j in range(count)]
            t=doc.add_table(rows=0,cols=count); t.autofit=False
            for col,w in zip(t.columns,widths): col.width=Cm(w)
            borders=OxmlElement('w:tblBorders')
            for edge in ('top','left','bottom','right','insideH','insideV'):
                el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            t._tbl.tblPr.append(borders)
            for j,row in enumerate(values):
                cells=t.add_row().cells
                trpr=t.rows[-1]._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
                if j==0: trpr.append(OxmlElement('w:tblHeader'))
                for k,(cell,value) in enumerate(zip(cells,row)):
                    cell.width=Cm(widths[k]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    pr=cell._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
                    for side in ('top','bottom','left','right'):
                        el=OxmlElement('w:'+side); el.set(qn('w:w'),'100'); el.set(qn('w:type'),'dxa'); margins.append(el)
                    pr.append(margins)
                    if j==0:
                        shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7EDF2'); pr.append(shade)
                    p=cell.paragraphs[0]; p.paragraph_format.line_spacing=1.15
                    p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3)
                    # 머리행이 쪽 끝에 홀로 남지 않도록 첫 데이터 행과 묶고, 4행 이하의 짧은 표는 통째로 묶는다
                    if j==0 or (len(values)<=5 and j<len(values)-1): p.paragraph_format.keep_with_next=True
                    if short_col[k]: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    inline(p,value)
                    for r in p.runs: helper.font_run(r,9); r.bold=j==0
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
        elif line.startswith('#'):
            heading=re.sub(r'^#+\s*','',line)
            # 절 번호의 점(1.1)과 문장부호는 남기고 Markdown 기호만 지운다
            m=re.match(r'^(\d+(?:\.\d+)*)\.?\s+(.*)$',heading)
            num,rest=(m.group(1),m.group(2)) if m else ('',heading)
            rest=re.sub(r'[*`\[\]#|<>]','',rest)
            heading=re.sub(r'\s+',' ',f'{num} {rest}').strip()
            doc.add_paragraph(heading,style='Heading 1' if line.startswith('## ') else 'Heading 2')
        else:
            while i<len(lines) and lines[i].strip() and not re.match(r'^(#|\||---|[-*] |\d+\. )',lines[i].strip()):
                line+=' '+lines[i].strip(); i+=1
            p=doc.add_paragraph()
            if line.startswith('- '): line='• '+line[2:]
            inline(p,line)
    if notes:
        doc.add_heading('자료 주석',level=1)
        doc.add_paragraph('자료 경로는 저장소의 docs/papers를 기준으로 한다. 원천 전체의 공개 접근이나 독립 재현을 보장하지 않는다.')
        for target,(num,label) in notes.items():
            p=doc.add_paragraph()
            el=OxmlElement('w:bookmarkStart'); el.set(qn('w:id'),str(num)); el.set(qn('w:name'),f'source{num}'); p._p.append(el)
            r=p.add_run(f'{num}. {label}'); helper.font_run(r,9); r.bold=True
            el=OxmlElement('w:bookmarkEnd'); el.set(qn('w:id'),str(num)); p._p.append(el)
            r=p.add_run('\n'+target); helper.font_run(r,9)
            p.paragraph_format.line_spacing=1.15
    out=out_dir/f'PlaceOS_{track}_논문_{stamp}.docx'
    out.parent.mkdir(parents=True,exist_ok=True)
    doc.save(out); helper.embed_fonts(out)
    return {'track':track,'path':str(out),'paragraphs':len(doc.paragraphs),'tables':table_count,'notes':len(notes)}


if __name__=='__main__':
    ap=argparse.ArgumentParser(description=__doc__,formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--src',type=Path,default=None)
    ap.add_argument('--out',type=Path,default=BASE/'output')
    ap.add_argument('--stamp',default='20260919')
    a=ap.parse_args()
    src_dir=(a.src or default_src()).resolve()
    print(json.dumps([build(k,v,src_dir,a.out.resolve(),a.stamp) for k,v in TITLES.items()],ensure_ascii=False,indent=2))
